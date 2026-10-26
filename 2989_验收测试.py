# SPDX-License-Identifier: MIT
"""2989：逐来源独立酒店月表业务验收；不使用任何库创建 XLSX。"""
from __future__ import annotations

import copy
import argparse
import csv
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import importlib.util
import json
import math
import os
import platform
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

BASE = Path(__file__).resolve().parent
HEADER = ["Rezervasyon No", "Ad", "Soyad", "Giriş Tarihi", "Çıkış Tarihi", "Ücret", "Notlar"]
SOURCE = BASE / "2989_自造月表.xlsx"
OUTPUTS = {"2989_处理模型.json", "2989_有效预订.csv", "2989_来源记录.csv", "2989_期间汇总.csv", "2989_处理报告.json"}


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, BASE / filename)
    loaded = importlib.util.module_from_spec(spec)
    sys.modules[name] = loaded
    spec.loader.exec_module(loaded)
    return loaded


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Acceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = module("engine2989_tests", "2989_合并引擎.py")
        cls.fixture = module("fixture2989_tests", "2989_演示数据.py")
        cls.model = cls.fixture.build_model()

    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="2989_验收暂存_"))
        self.output = self.root / "2989_处理结果"

    def tearDown(self):
        checked = self.root.resolve()
        if checked.parent != Path(tempfile.gettempdir()).resolve() or not checked.name.startswith("2989_验收暂存_"):
            raise RuntimeError("测试目录不在拥有范围，拒绝清理")
        shutil.rmtree(checked)

    def row(self, booking="2989_A", first="İPEK", last="YILMAZ", arrival="2023-01-15", departure="2023-01-17", price="USD 80", notes="2989 demo"):
        return [booking, first, last, arrival, departure, price, notes]

    def analyze(self, rows, header=None, bare_currency="USD", name="2989_测试月表"):
        return self.engine.analyze_sheets([dict(name=name, rows=[header or HEADER, *rows])], bare_currency=bare_currency)

    def counts(self, model, accepted, review=0, duplicate=0):
        counts = model["report"]["counts"]
        self.assertEqual(counts["accepted"], accepted)
        self.assertEqual(counts["review"], review)
        self.assertEqual(counts["duplicate"], duplicate)
        self.assertEqual(len(model["accepted"]), accepted)
        self.assertEqual(len(model["review"]), review)
        self.assertEqual(sum(row["status"] == "duplicate" for row in model["ledger"]), duplicate)

    def source_required(self):
        self.assertTrue(SOURCE.is_file(), "须先用项目JS构建器生成真实自造XLSX；测试不会代写XLSX")

    def test_01_full_48_sheet_json_model_all_1540_source_statuses(self):
        untouched = copy.deepcopy(self.model["sheets"])
        model = self.engine.analyze_sheets(self.model["sheets"], bare_currency="USD")
        self.assertEqual(self.model["sheets"], untouched)
        result = self.fixture.verify_model(model, self.model)
        self.assertEqual(result["verified_accepted_records"], 1430)
        self.assertEqual(result["total_revenue_usd"], 171600)

    def test_02_every_physical_row_has_exactly_one_source_ledger_record(self):
        model = self.engine.analyze_sheets(self.model["sheets"], bare_currency="USD")
        expected = {(s["name"], i + 1) for s in self.model["sheets"] for i in range(len(s["rows"]))}
        actual = [(r["source_sheet"], r["source_row"]) for r in model["ledger"]]
        self.assertEqual(len(actual), len(set(actual)))
        self.assertEqual(set(actual), expected)
        self.assertTrue(all(r["reason"] and r["status"] for r in model["ledger"]))
        for record in model["ledger"]:
            original = next(s for s in self.model["sheets"] if s["name"] == record["source_sheet"])["rows"][record["source_row"] - 1]
            self.assertEqual(json.loads(record["raw_json"]), original)

    def test_03_known_aliases_reorder_and_unicode_header_whitespace(self):
        header = ["  ÜCRET  ", "ÇIKIŞ   TARİHİ", "ADı", "REZ.NO", "GİRİŞ TARİHİ", "SOYADı", "AÇIKLAMA"]
        model = self.analyze([["USD 80", "2023-01-17", "İPEK", "2989_A", "2023-01-15", "YILMAZ", "ok"]], header)
        self.counts(model, 1)
        self.assertEqual(model["accepted"][0]["first_name"], "İpek")
        self.assertEqual(model["accepted"][0]["last_name"], "Yılmaz")

    def test_04_turkish_title_preserves_dotted_and_dotless_letters(self):
        for raw, expected in [("İPEK", "İpek"), ("YILMAZ", "Yılmaz"), ("IŞIK", "Işık"), ("şahin", "Şahin"),
                              ("  İPEK   YILMAZ  ", "İpek Yılmaz"), ("sahin", "Sahin"), ("Yilmaz", "Yilmaz"),
                              ("AYŞE-İPEK", "Ayşe-İpek"), ("o'neil", "O'Neil")]:
            with self.subTest(raw=raw):
                self.assertEqual(self.engine.turkish_title(raw), expected)
        model = self.analyze([self.row(first="ŞAHİN", last="İPEK")])
        self.assertEqual((model["accepted"][0]["first_name"], model["accepted"][0]["last_name"]), ("Şahin", "İpek"))
        for raw in ["---", "'''", "’ ’"]:
            with self.subTest(punctuation_only_name=raw):
                self.counts(self.analyze([self.row(first=raw)]), 0, review=1)

    def test_05_explicit_and_native_dates_are_exact(self):
        for value in [date(2024, 2, 29), datetime(2024, 2, 29), {"$date": "2024-02-29"}, "2024-02-29", "29.02.2024", "2024/02/29"]:
            with self.subTest(value=value):
                self.assertEqual(self.engine.parse_date(value), date(2024, 2, 29))

    def test_06_ambiguous_day_month_dates_never_guessed(self):
        for value in ["01/02/2024", "01.02.2024", "02-01-2024", "12/11/2024"]:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    self.engine.parse_date(value)

    def test_07_invalid_and_unmarked_numeric_dates_are_rejected(self):
        values = ["2023-02-29", "2024-13-01", "2024-01-32", "", None, 45200, True,
                  datetime(2024, 1, 1, 1), datetime(2024, 1, 1, tzinfo=timezone.utc)]
        for value in values:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    self.engine.parse_date(value)

    def test_08_explicit_usd_price_styles_and_bare_opt_in(self):
        for value in ["USD 80", "80 USD", "$80.00", "USD 80,00", "80,00 USD"]:
            with self.subTest(value=value):
                self.assertEqual(self.engine.parse_price(value), Decimal(80))
        for value in [80, 80.0, Decimal("80.00"), "80.00"]:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    self.engine.parse_price(value)
                self.assertEqual(self.engine.parse_price(value, bare_currency="USD"), Decimal(80))

    def test_09_default_bare_currency_is_review_instead_of_assumed(self):
        model = self.analyze([self.row(price=80), self.row("2989_B", price="USD 80")], bare_currency=None)
        self.counts(model, 1, review=1)
        self.assertEqual(model["accepted"][0]["booking_id"], "2989_B")

    def test_10_negative_non_usd_nonfinite_and_fractional_cent_prices_reject(self):
        for value in ["TRY 80", "80 TL", "₺80", "EUR 80", "USD -80", -80, "USD 80.001", True, float("nan"), float("inf")]:
            with self.subTest(value=repr(value)):
                with self.assertRaises(ValueError):
                    self.engine.parse_price(value, bare_currency="USD")

    def test_11_same_id_equivalent_records_keep_earliest_source(self):
        rows = [self.row(), self.row(first=" ipek ", last=" yılmaz ", arrival={"$date": "2023-01-15"}, departure="17.01.2023", price="$80.00")]
        model = self.analyze(rows)
        self.counts(model, 1, duplicate=1)
        self.assertEqual(model["accepted"][0]["source_row"], 2)
        self.assertEqual(model["ledger"][2]["status"], "duplicate")
        header = HEADER + ["2989 Room"]
        for room2, expected_accepted, expected_review, expected_duplicate in [("101", 1, 0, 1), ("102", 0, 2, 0)]:
            with self.subTest(same_id_extra_room=room2):
                model = self.analyze([self.row() + ["101"], self.row() + [room2]], header)
                self.counts(model, expected_accepted, expected_review, expected_duplicate)
                if model["accepted"]:
                    self.assertNotIn("_extra_signature", model["accepted"][0])

    def test_12_same_id_conflict_or_bad_member_reviews_entire_group(self):
        sets = [[self.row(), self.row(departure="2023-01-18", price="USD 120")],
                [self.row(), self.row(departure="")]]
        for rows in sets:
            for order in [rows, list(reversed(rows))]:
                with self.subTest(rows=order):
                    self.counts(self.analyze(order), 0, review=2)
        before_header = [self.row(departure=""), HEADER, self.row()]
        model = self.engine.analyze_sheets([dict(name="2989_表头前坏行", rows=before_header)], bare_currency="USD")
        self.counts(model, 0, review=2)
        self.assertEqual([row["status"] for row in model["ledger"]], ["review", "header", "review"])

    def test_13_missing_id_unique_is_valid_but_same_identity_all_review(self):
        self.counts(self.analyze([self.row(booking="")]), 1)
        self.counts(self.analyze([self.row(booking=""), self.row(booking="", first="ipek", price="$80")]), 0, review=2)
        model = self.analyze([self.row(booking="") + ["101"], self.row(booking="") + ["102"]], HEADER + ["2989 Room"])
        self.counts(model, 0, review=2)

    def test_14_nonpositive_nights_missing_fields_and_rate_mismatch_review(self):
        rows = [self.row("2989_A", departure="2023-01-15"), self.row("2989_B", departure="2023-01-14"),
                self.row("2989_C", price="USD 120"), self.row("2989_D", departure=""), self.row("2989_E", last="")]
        self.counts(self.analyze(rows), 0, review=5)

    def test_15_unknown_and_duplicate_alias_headers_do_not_silently_skip_rows(self):
        for header in [["Unknown", "Customer", "Date"], HEADER + ["İsim"]]:
            with self.subTest(header=header):
                model = self.analyze([self.row()], header)
                self.assertEqual(len(model["ledger"]), 2)
                self.assertTrue(all(row["status"] == "review" for row in model["ledger"]))

    def test_16_later_changed_header_mapping_reviews_whole_sheet(self):
        rows = [self.row(), list(reversed(HEADER)), list(reversed(self.row("2989_B")))]
        model = self.analyze(rows)
        self.counts(model, 0, review=4)

    def test_17_notes_titles_blanks_headers_and_totals_are_accounted(self):
        rows = [["2989 synthetic title"], HEADER, self.row(), [], HEADER, ["TOPLAM", 80]]
        model = self.engine.analyze_sheets([dict(name="2989_酒店", rows=rows), dict(name="2989_NOTLAR", rows=[["NOTLAR"], [], ["2989 notes"]])], bare_currency="USD")
        self.counts(model, 1)
        counts = model["report"]["counts"]
        self.assertEqual(counts["metadata"], 3)
        self.assertEqual(counts["blank"], 2)
        self.assertEqual(counts["header"], 1)
        self.assertEqual(counts["repeated_header"], 1)
        self.assertEqual(counts["total"], 1)
        self.assertEqual(counts["physical_rows"], 9)
        for unknown_rows in [[["NOTLAR"], ["2989 note", 80]], [["NOTLAR"], [80]]]:
            with self.subTest(notlar_with_possible_data=unknown_rows):
                model = self.engine.analyze_sheets([dict(name="2989_不明说明页", rows=unknown_rows)])
                self.counts(model, 0, review=2)

    def test_18_real_booking_named_total_is_not_discarded(self):
        for booking, first in [("TOTAL", "İPEK"), ("2989_A", "Total")]:
            with self.subTest(booking=booking, first=first):
                self.counts(self.analyze([self.row(booking=booking, first=first)]), 1)
        for arrival, departure in [("2023-01-15", "2023-01-17"), ("15.01.2023", "17.01.2023")]:
            with self.subTest(total_id_missing_names_with_date_columns=arrival):
                model = self.analyze([self.row(booking="TOTAL", first="", last="", arrival=arrival, departure=departure, price=80)])
                self.counts(model, 0, review=1)
                self.assertEqual(model["report"]["counts"]["total"], 0)

    def test_19_formula_string_is_audited_and_never_evaluated(self):
        original = self.row(price="=40*2", notes='=HYPERLINK("https://example.invalid")')
        model = self.analyze([original])
        self.counts(model, 0, review=1)
        self.assertEqual(json.loads(model["review"][0]["raw_json"]), original)
        self.assertIn("公式", model["review"][0]["reason"])

    def test_20_optional_identifier_and_notes_extra_column_preservation(self):
        header = ["Ad", "Soyad", "Giriş", "Çıkış", "Fiyat", "2989 extra"]
        original = ["İPEK", "YILMAZ", "2023-01-15", "2023-01-17", "$80", "raw extra"]
        model = self.analyze([original], header)
        self.counts(model, 1)
        self.assertEqual(model["accepted"][0]["booking_id"], "")
        self.assertEqual(model["accepted"][0]["notes"], "")
        self.assertEqual(json.loads(model["ledger"][1]["raw_json"])[-1], "raw extra")

    def test_21_cross_month_full_price_assigned_to_checkin_month(self):
        model = self.analyze([self.row(arrival="2023-01-31", departure="2023-02-02")])
        summaries = {r["period"]: r for r in model["summary"]}
        self.assertEqual((summaries["2023-01"]["bookings"], summaries["2023-01"]["revenue_usd"]), (1, 80))
        self.assertEqual((summaries["2023-02"]["bookings"], summaries["2023-02"]["revenue_usd"]), (0, 0))
        self.assertEqual(summaries["2023"]["revenue_usd"], 80)

    def test_22_global_shape_and_capacity_limits_reject(self):
        for sheets in [[], [dict(name="2989_A", rows=[])] * 101, [dict(name="2989_A", rows=[]), dict(name="2989_A", rows=[])]]:
            with self.subTest(count=len(sheets)):
                with self.assertRaises(ValueError):
                    self.engine.analyze_sheets(sheets)
        with patch.object(self.engine, "MAX_ROWS", 1):
            with self.assertRaises(ValueError):
                self.analyze([self.row()])
        with patch.object(self.engine, "MAX_CELLS", 10):
            with self.assertRaises(ValueError):
                self.analyze([self.row()])

    def test_23_real_sdk_xlsx_full_oracle_native_dates_and_source_hash(self):
        self.source_required()
        import openpyxl
        original = digest(SOURCE)
        workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False, keep_links=False)
        try:
            self.assertEqual(len(workbook.worksheets), 48)
            self.assertTrue(any(s.max_row is None or s.max_column is None for s in workbook),
                            "自造SDK源应真实覆盖缺少可选dimension的兼容回归")
            native_dates = sum(isinstance(cell.value, (date, datetime)) for sheet in workbook for row in sheet for cell in row)
            self.assertGreater(native_dates, 0)
        finally:
            workbook.close()
        report = self.engine.process_workbook(SOURCE, self.output, bare_currency="USD", export_xlsx=False)
        self.assertEqual(set(p.name for p in self.output.iterdir()), OUTPUTS)
        model = json.loads((self.output / "2989_处理模型.json").read_text(encoding="utf-8"))
        self.fixture.verify_model(model, self.model)
        self.assertEqual(report["source_sha256"], original)
        self.assertEqual(digest(SOURCE), original)
        with (self.output / "2989_有效预订.csv").open(encoding="utf-8-sig", newline="") as stream:
            records = list(csv.DictReader(stream))
        self.assertEqual(len(records), 1430)
        self.assertEqual(sum(Decimal(r["price_usd"]) for r in records), Decimal(171600))
        with (self.output / "2989_来源记录.csv").open(encoding="utf-8-sig", newline="") as stream:
            source_rows = [json.loads(r["source_record_json"]) for r in csv.DictReader(stream)]
        self.assertEqual(source_rows, model["ledger"])

    def test_24_existing_empty_and_occupied_outputs_never_changed(self):
        self.source_required()
        before = digest(SOURCE)
        self.output.mkdir()
        with self.assertRaises(ValueError):
            self.engine.process_workbook(SOURCE, self.output, bare_currency="USD")
        marker = self.output / "2989_已有记录.txt"
        marker.write_bytes(b"2989 keep")
        with self.assertRaises(ValueError):
            self.engine.process_workbook(SOURCE, self.output, bare_currency="USD")
        self.assertEqual(marker.read_bytes(), b"2989 keep")
        self.assertEqual(digest(SOURCE), before)

    def test_25_write_failure_and_external_source_change_do_not_publish(self):
        self.source_required()
        source = self.root / "2989_副本.xlsx"
        source.write_bytes(SOURCE.read_bytes())
        before = set(self.root.iterdir())
        with patch.object(csv, "writer", side_effect=OSError("2989 simulated storage error")):
            with self.assertRaises(OSError):
                self.engine.process_workbook(source, self.output, bare_currency="USD")
        self.assertFalse(self.output.exists())
        self.assertEqual(set(self.root.iterdir()), before)
        original_analyze = self.engine.analyze_sheets
        def mutate(*args, **kwargs):
            result = original_analyze(*args, **kwargs)
            source.write_bytes(source.read_bytes() + b"2989 changed externally")
            return result
        with patch.object(self.engine, "analyze_sheets", side_effect=mutate):
            with self.assertRaises(ValueError):
                self.engine.process_workbook(source, self.output, bare_currency="USD")
        self.assertFalse(self.output.exists())
        self.assertEqual(set(self.root.iterdir()), before)

    def test_26_unsupported_extensions_macro_and_external_link_preflight(self):
        self.source_required()
        for suffix in [".xls", ".xlsm"]:
            source = self.root / ("2989_不支持" + suffix)
            source.write_bytes(SOURCE.read_bytes())
            with self.subTest(suffix=suffix):
                with self.assertRaises(ValueError):
                    self.engine.process_workbook(source, self.output, bare_currency="USD")
                self.assertFalse(self.output.exists())
        real_infos = zipfile.ZipFile.infolist
        for name in ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"]:
            def injected(package):
                return real_infos(package) + [zipfile.ZipInfo(name)]
            with self.subTest(zip_member=name), patch.object(zipfile.ZipFile, "infolist", injected):
                with self.assertRaises(ValueError):
                    self.engine.process_workbook(SOURCE, self.output, bare_currency="USD")
                self.assertFalse(self.output.exists())

    def test_27_actual_windows_junction_input_and_output_paths_rejected(self):
        self.source_required()
        if os.name != "nt":
            self.skipTest("此项验证Windows目录连接点")
        real = self.root / "2989_真实目录"
        real.mkdir()
        copied = real / "2989_副本.xlsx"
        copied.write_bytes(SOURCE.read_bytes())
        junction = self.root / "2989_连接入口"
        created = subprocess.run(["cmd", "/c", "mklink", "/J", str(junction), str(real)], capture_output=True)
        if created.returncode:
            self.skipTest("系统未允许创建目录连接点")
        try:
            for source, output in [(junction / copied.name, self.output), (copied, junction / "2989_新结果")]:
                with self.subTest(source=str(source), output=str(output)):
                    with self.assertRaises(ValueError):
                        self.engine.process_workbook(source, output, bare_currency="USD")
                    self.assertFalse(output.exists())
        finally:
            if junction.parent != self.root or not junction.name.startswith("2989_"):
                raise RuntimeError("连接点不在测试拥有范围")
            junction.rmdir()


class TimedResult(unittest.TextTestResult):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.details = []

    def startTest(self, test):
        super().startTest(test)
        self.current = dict(test=test._testMethodName, status="running")
        self.current_started = time.perf_counter()

    def addSuccess(self, test):
        super().addSuccess(test)
        self.current["status"] = "passed"

    def addFailure(self, test, error):
        super().addFailure(test, error)
        self.current["status"] = "failure"

    def addError(self, test, error):
        super().addError(test, error)
        self.current["status"] = "error"

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self.current.update(status="skipped", reason=reason)

    def addSubTest(self, test, subtest, error):
        super().addSubTest(test, subtest, error)
        if error:
            self.current["status"] = "failure" if issubclass(error[0], test.failureException) else "error"

    def stopTest(self, test):
        self.current["elapsed_seconds"] = round(time.perf_counter() - self.current_started, 6)
        self.details.append(self.current)
        super().stopTest(test)


def main():
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--report", type=Path, help="可选：项目根目录下2989前缀JSON真实测试记录")
    args, remaining = parser.parse_known_args()
    report_path = args.report
    if report_path is not None:
        report_path = Path(os.path.abspath(report_path))
        if report_path.parent != BASE or not report_path.name.startswith("2989_") or report_path.suffix != ".json" or report_path.is_symlink():
            raise ValueError("测试报告须为本项目根目录内的2989前缀JSON普通文件")
    names = ["2989_合并引擎.py", "2989_演示数据.py", "2989_验收测试.py"]
    before = {name: digest(BASE / name) for name in names}
    started = time.perf_counter()
    runner = unittest.TextTestRunner(verbosity=2, resultclass=TimedResult)
    program = unittest.main(argv=[sys.argv[0], *remaining], testRunner=runner, exit=False)
    result = program.result
    independent = None
    if result.wasSuccessful() and SOURCE.is_file():
        analyzed = Acceptance.engine.analyze_sheets(Acceptance.engine.read_workbook(SOURCE), bare_currency="USD", source_sha256=digest(SOURCE))
        independent = Acceptance.fixture.verify_model(analyzed, Acceptance.model)
        independent["source_xlsx_sha256"] = digest(SOURCE)
    after = {name: digest(BASE / name) for name in names}
    if report_path is not None:
        report = dict(schema_version=1, run_at_utc=datetime.now(timezone.utc).isoformat(),
                      command="python -B 2989_验收测试.py --report 2989_测试结果.json", tests_run=result.testsRun,
                      passed=sum(r["status"] == "passed" for r in result.details),
                      failures=sum(r["status"] == "failure" for r in result.details), errors=sum(r["status"] == "error" for r in result.details),
                      skipped=len(result.skipped), successful=result.wasSuccessful() and before == after,
                      elapsed_seconds=round(time.perf_counter() - started, 6), per_test=result.details,
                      environment=dict(python=platform.python_version(), implementation=platform.python_implementation(), system=platform.system(), release=platform.release()),
                      code_sha256=after, code_unchanged_during_tests=before == after,
                      source_xlsx_sha256=digest(SOURCE) if SOURCE.is_file() else None,
                      independent_source_model_verification=independent,
                      workbook_authorship="Tests do not create XLSX. Actual workbook authored by the separate project JS builder; openpyxl used read-only.",
                      expected_demo_counts=dict(sheets=48, input_order_rows=1540, accepted=1430, review=80, duplicate=30, revenue_usd=171600))
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    raise SystemExit(0 if result.wasSuccessful() and before == after else 1)


if __name__ == "__main__":
    main()
