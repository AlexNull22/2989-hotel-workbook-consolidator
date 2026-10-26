# SPDX-License-Identifier: MIT
"""只读提取异构 XLSX，保守规范化与可追溯分流。"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

PREFIX = "2989"
RATE = Decimal("40")
ALIASES = {
    "booking_id": ["Rezervasyon No", "Rez.No", "Booking No"],
    "first_name": ["Ad", "Adı", "İsim", "First Name"],
    "last_name": ["Soyad", "Soyadı", "Last Name"],
    "check_in": ["Giriş", "Giriş Tarihi", "Check-in"],
    "check_out": ["Çıkış", "Çıkış Tarihi", "Check-out"],
    "price": ["Fiyat", "Ücret", "Tutar", "Price"],
    "notes": ["Not", "Notlar", "Açıklama", "Notes"],
}
REQUIRED = {"first_name", "last_name", "check_in", "check_out", "price"}
MAX_ROWS = 100_000
MAX_CELLS = 2_000_000


def tr_lower(value: str) -> str:
    return unicodedata.normalize("NFKC", value).replace("I", "ı").replace("İ", "i").lower()


def key(value) -> str:
    return " ".join(tr_lower(str(value or "")).split())


LOOKUP = {key(alias): field for field, aliases in ALIASES.items() for alias in aliases}


def turkish_title(value: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("姓名缺失或不是文本")
    value = " ".join(unicodedata.normalize("NFKC", value).split())
    if len(value) > 100 or not any(c.isalpha() for c in value) or not all(c.isalpha() or c in " '-’" for c in value):
        raise ValueError("姓名包含不支持的字符")
    parts = re.split(r"([ '\-’])", tr_lower(value))
    return "".join((p[0].replace("i", "İ").replace("ı", "I").upper() + p[1:])
                   if p and p not in " '-’" else p for p in parts)


def parse_date(value) -> date:
    if isinstance(value, dict) and set(value) == {"$date"}:
        value = datetime.fromisoformat(value["$date"])
    if isinstance(value, datetime):
        if value.tzinfo is not None or value.time() != datetime.min.time():
            raise ValueError("日期带有时间或时区，需复核")
        answer = value.date()
    elif isinstance(value, date):
        answer = value
    elif isinstance(value, str):
        value = value.strip()
        if re.fullmatch(r"\d{4}[-/]\d{2}[-/]\d{2}", value):
            answer = date.fromisoformat(value.replace("/", "-"))
        else:
            found = re.fullmatch(r"(\d{1,2})([./-])(\d{1,2})\2(\d{4})", value)
            if not found:
                raise ValueError("日期格式不支持或缺失")
            a, b, year = int(found[1]), int(found[3]), int(found[4])
            candidates = set()
            for month, day in [(b, a), (a, b)]:
                try:
                    candidates.add(date(year, month, day))
                except ValueError:
                    pass
            if len(candidates) != 1:
                raise ValueError("日期有日/月歧义或无效")
            answer = candidates.pop()
    else:
        raise ValueError("日期缺失或未经日期格式标记的数字")
    if not 1900 <= answer.year <= 2100:
        raise ValueError("日期超出1900至2100年范围")
    return answer


def parse_price(value, bare_currency=None) -> Decimal:
    if bare_currency not in (None, "USD"):
        raise ValueError("裸金额币种选项仅支持USD或未指定")
    if isinstance(value, bool) or value is None:
        raise ValueError("金额缺失或类型无效")
    explicit = False
    if isinstance(value, str):
        text = value.strip()
        if re.search(r"TRY|TL|₺|EUR|€|GBP|£", text, re.I):
            raise ValueError("非USD金额，不自动换汇")
        matched = re.fullmatch(r"(?:USD\s*|\$\s*)?([+-]?\d+(?:[.,]\d{1,2})?)(?:\s*USD)?", text, re.I)
        if not matched:
            raise ValueError("金额格式不支持，千分位或文本需复核")
        explicit = bool(re.search(r"USD|\$", text, re.I))
        text = matched[1].replace(",", ".")
    elif isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        raise ValueError("金额类型无效")
    if not explicit and bare_currency != "USD":
        raise ValueError("裸金额币种未指定")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("金额无效") from exc
    if not amount.is_finite() or not 0 <= amount <= Decimal("146000") or amount != amount.quantize(Decimal("0.01")):
        raise ValueError("金额为负、超过上限或超过两位小数")
    return amount


def json_default(value):
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(type(value).__name__)


def raw_json(row):
    return json.dumps(row, ensure_ascii=False, separators=(",", ":"), default=json_default, allow_nan=False)


def header_mapping(row):
    result = {}
    duplicate = False
    for i, value in enumerate(row):
        field = LOOKUP.get(key(value))
        if field:
            if field in result:
                duplicate = True
            result[field] = i
    return result if REQUIRED <= result.keys() and not duplicate else None


def total_row(row):
    nonempty=[(i,v) for i,v in enumerate(row) if v is not None and v!=""]
    if not nonempty or nonempty[0][0]>1 or key(nonempty[0][1]) not in {"toplam","genel toplam","total"}:
        return False
    # 只接受单个明确标记加数字/金额，不能把名为Total的真实预订误删。
    return all(isinstance(v,(int,float,Decimal)) and not isinstance(v,bool)
               or isinstance(v,str) and re.fullmatch(r"[+-]?(?:[$₺]\s*)?\d+(?:[.,]\d{1,2})?",v.strip())
               for _,v in nonempty[1:])


def normalized_record(row, mapping):
    def get(field):
        pos = mapping.get(field)
        return row[pos] if pos is not None and pos < len(row) else None
    reasons = []
    ident = get("booking_id")
    if ident is None or ident == "":
        ident = ""
    elif not isinstance(ident, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", ident.strip()):
        reasons.append("订单号应为不超过64字符的文本标识")
        ident = ""
    else:
        ident = ident.strip()
    result = {"booking_id": ident}
    for field in ("first_name", "last_name"):
        try:
            result[field] = turkish_title(get(field))
        except ValueError as exc:
            reasons.append(f"{field}: {exc}")
    for field in ("check_in", "check_out"):
        try:
            result[field] = parse_date(get(field))
        except (ValueError, TypeError) as exc:
            reasons.append(f"{field}: {exc}")
    notes = get("notes")
    if notes is None:
        notes = ""
    if not isinstance(notes, str) or len(notes) > 2000:
        reasons.append("备注应为不超过2000字符的文本")
        notes = ""
    result["notes"] = notes.strip()
    return result, get("price"), reasons


def analyze_sheets(sheets, bare_currency=None, source_sha256="synthetic"):
    if bare_currency not in (None, "USD"):
        raise ValueError("裸金额币种选项无效")
    if not 1 <= len(sheets) <= 100:
        raise ValueError("工作表数应为1至100")
    if len({s["name"] for s in sheets}) != len(sheets):
        raise ValueError("工作表名重复")
    if sum(len(s["rows"]) for s in sheets) > MAX_ROWS:
        raise ValueError("工作簿超过100000行")
    if sum(len(r) for s in sheets for r in s["rows"]) > MAX_CELLS:
        raise ValueError("工作簿超过2000000单元格")
    ledger, candidates, sheet_details = [], [], []
    for sheet in sheets:
        rows, name = sheet["rows"], sheet["name"]
        if any(len(str(v)) > 10000 for row in rows for v in row if v is not None):
            raise ValueError("单元格超过10000字符")
        heads = [(i, header_mapping(row)) for i, row in enumerate(rows[:20]) if header_mapping(row)]
        note_only = bool(rows and rows[0] and key(rows[0][0]) == "notlar" and not heads
                         and all(len([v for v in row if v is not None and v!=""])<=1
                                 and all(v is None or isinstance(v,str) for v in row) for row in rows))
        mapping = heads[0][1] if heads else None
        # 发现任何后续完整表头必须保持同一映射，不能悄悄错列。
        conflict = bool(mapping and any(header_mapping(row) not in (None, mapping) for row in rows))
        first_head = heads[0][0] if heads else -1
        for i, row in enumerate(rows):
            pre_header_reason = []
            item = {"source_sheet": name, "source_row": i+1, "status": "", "reason": "", "raw_json": raw_json(row)}
            ledger.append(item)
            if not any(v is not None and v != "" for v in row):
                item.update(status="blank", reason="空白行")
                continue
            if note_only:
                item.update(status="metadata", reason="NOTLAR明确说明页")
                continue
            if not mapping or conflict:
                item.update(status="review", reason="未知/重复别名表头" if not mapping else "多个表头映射冲突")
                continue
            if i < first_head:
                nonempty=[v for v in row if v is not None and v!=""]
                title_only=len(nonempty)==1 and isinstance(nonempty[0],str)
                if title_only:
                    item.update(status="metadata", reason="已识别表头之前的单格标题")
                    continue
                pre_header_reason.append("表头前存在多格内容，不能确认是标题")
            if header_mapping(row) == mapping:
                item.update(status="header" if i == first_head else "repeated_header", reason="字段表头")
                continue
            # 仅独占首个非空单元格为已知总计标记时判为总计；姓名相似不误删。
            if not pre_header_reason and total_row(row):
                item.update(status="total", reason="明确合计标记，未作为预订收入")
                continue
            record, price, reasons = normalized_record(row, mapping)
            reasons.extend(pre_header_reason)
            extra_fields = [[key(rows[first_head][j]) if j<len(rows[first_head]) else f"column_{j+1}",v]
                            for j,v in enumerate(row) if j not in mapping.values() and v is not None and v!=""]
            record["_extra_signature"] = raw_json(extra_fields)
            if any(isinstance(v,str) and v.startswith("=") for v in row):
                reasons.append("源行含公式，未执行")
            try:
                record["price_usd"] = parse_price(price, bare_currency)
            except ValueError as exc:
                reasons.append(str(exc))
            if "check_in" in record and "check_out" in record:
                nights = (record["check_out"] - record["check_in"]).days
                record["nights"] = nights
                if not 1 <= nights <= 3650:
                    reasons.append("退房未晚于入住或超过3650晚")
                elif "price_usd" in record and record["price_usd"] != RATE * nights:
                    reasons.append("金额与晚数×40美元不一致")
            candidates.append((item,record,reasons))
        sheet_details.append({"sheet": name, "physical_rows": len(rows), "header_row": first_head+1 if heads else None, "recognized": bool(mapping and not conflict), "note_only": note_only})

    by_id = defaultdict(list)
    no_id = []
    fields = ("first_name", "last_name", "check_in", "check_out", "nights", "price_usd", "notes")
    for entry in candidates:
        if entry[1]["booking_id"]:
            by_id[entry[1]["booking_id"]].append(entry)
        else:
            no_id.append(entry)
    accepted = []

    def signature(record, include_extras=True):
        return tuple(record.get(k) for k in fields) + ((record.get("_extra_signature"),) if include_extras else ())

    def retain(entry):
        item, record, _ = entry
        item.update(status="accepted", reason="通过明确结构、日期、币种与固定房价核对")
        cooked = dict(record)
        cooked.pop("_extra_signature",None)
        for k in ("check_in", "check_out"):
            cooked[k] = cooked[k].isoformat()
        cooked["price_usd"] = int(cooked["price_usd"])  # 固定40美元整夜计价，无小数截断。
        cooked.update(source_sheet=item["source_sheet"],source_row=item["source_row"])
        accepted.append(cooked)

    for group in by_id.values():
        if any(e[2] for e in group) or len({signature(e[1]) for e in group}) > 1:
            for item, _, reasons in group:
                item.update(status="review", reason="；".join(reasons + (["同订单号整组存在冲突或坏行"] if len(group)>1 else [])))
        else:
            retain(group[0])
            for item, _, _ in group[1:]:
                item.update(status="duplicate",reason="同订单号规范字段一致；保留来源最早行")
    no_id_groups = defaultdict(list)
    for entry in no_id:
        no_id_groups[signature(entry[1],include_extras=False)].append(entry)
    for group in no_id_groups.values():
        for entry in group:
            item, _, reasons = entry
            if len(group)>1 or reasons:
                item.update(status="review",reason="；".join(reasons + (["缺订单号且身份字段重复，不能确认是否多房预订"] if len(group)>1 else [])))
            else:
                retain(entry)
    accepted.sort(key=lambda r:(r["check_in"],r["booking_id"],r["source_sheet"],r["source_row"]))
    monthly, yearly = defaultdict(lambda:[0,0]), defaultdict(lambda:[0,0])
    for r in accepted:
        for target, period in ((monthly,r["check_in"][:7]),(yearly,r["check_in"][:4])):
            target[period][0]+=1
            target[period][1]+=r["price_usd"]
    summary = []
    if accepted:
        first_year, last_year = int(accepted[0]["check_in"][:4]), int(accepted[-1]["check_in"][:4])
        for year in range(first_year,last_year+1):
            count,amount=yearly[str(year)]
            summary.append({"period":str(year),"bookings":count,"revenue_usd":amount})
        for year in range(first_year,last_year+1):
            for month in range(1,13):
                period=f"{year}-{month:02}"
                count,amount=monthly[period]
                summary.append({"period":period,"bookings":count,"revenue_usd":amount})
    counts = dict(Counter(row["status"] for row in ledger))
    for state in ("accepted","review","duplicate","blank","metadata","header","repeated_header","total"):
        counts.setdefault(state,0)
    counts.update(physical_rows=len(ledger),sheets=len(sheets),booking_rows=len(candidates),revenue_usd=sum(r["price_usd"] for r in accepted))
    return {"accepted":accepted,"ledger":ledger,"review":[r for r in ledger if r["status"]=="review"],"summary":summary,
            "report":{"source_sha256":source_sha256,"counts":counts,"sheet_details":sheet_details,"bare_currency":bare_currency,"nightly_rate_usd":40,"monthly_basis":"入住月归集整笔已核实预订金额，不是逐夜收入确认"}}


def safe_path(path):
    path=Path(os.path.abspath(path))
    for ancestor in (path,*path.parents):
        if ancestor.exists() and (ancestor.is_symlink() or (hasattr(ancestor,"is_junction") and ancestor.is_junction())):
            raise ValueError("路径含符号链接或目录连接点")
    return path


def read_workbook(source):
    import openpyxl  # 仅解析读取；所有XLSX写入由独立JS导出器负责。
    if source.suffix.lower() != ".xlsx" or not source.is_file() or source.stat().st_size>25_000_000:
        raise ValueError("仅支持不超过25MB的普通xlsx")
    with zipfile.ZipFile(source) as package:
        infos=package.infolist()
        if sum(i.file_size for i in infos)>100_000_000 or len(infos)>2000:
            raise ValueError("工作簿解压大小或文件项超过上限")
        names=[i.filename.lower() for i in infos]
        if any("vba" in n or "externallinks/" in n for n in names):
            raise ValueError("不支持宏或外部工作簿链接")
    workbook=openpyxl.load_workbook(source,read_only=True,data_only=False,keep_links=False)
    try:
        if len(workbook.worksheets)>100:
            raise ValueError("工作表超过100张")
        sheets=[]
        total_rows=0
        total_cells=0
        for sheet in workbook.worksheets:
            if (sheet.max_row is not None and sheet.max_row>MAX_ROWS) or (sheet.max_column is not None and sheet.max_column>100):
                raise ValueError("工作表尺寸超过100000行或100列")
            rows=[]
            for cells in sheet.iter_rows():
                row=[c.value for c in cells]
                while row and row[-1] is None:
                    row.pop()
                if len(row)>100:
                    raise ValueError("工作表超过100列")
                rows.append(row)
                total_rows+=1
                total_cells+=len(row)
                if total_rows>MAX_ROWS or total_cells>MAX_CELLS:
                    raise ValueError("工作簿行或单元格数量超过上限")
            sheets.append({"name":sheet.title,"rows":rows})
        return sheets
    finally:
        workbook.close()


def find_node():
    runtime=Path.home()/".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
    if runtime.is_file():
        return str(runtime)
    node=shutil.which("node")
    if not node:
        raise RuntimeError("未找到Node；可先用--no-xlsx仅导出CSV及JSON")
    return node


def csv_literal(value):
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value


def process_workbook(source,output_dir,bare_currency=None,export_xlsx=False):
    source=safe_path(source)
    output_dir=safe_path(output_dir)
    if output_dir.exists() or not output_dir.parent.is_dir():
        raise ValueError("结果目录必须不存在且父目录存在")
    if source.suffix.lower()!=".xlsx" or not source.is_file() or source.stat().st_size>25_000_000:
        raise ValueError("仅支持不超过25MB的普通xlsx")
    original=source.read_bytes()
    before=hashlib.sha256(original).hexdigest()
    model=analyze_sheets(read_workbook(source),bare_currency,before)
    model["report"]["source_name"]=source.name
    temp=Path(tempfile.mkdtemp(prefix="2989_暂存_",dir=output_dir.parent))
    try:
        model_path=temp/"2989_处理模型.json"
        model_path.write_text(json.dumps(model,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")
        with (temp/"2989_有效预订.csv").open("w",encoding="utf-8-sig",newline="") as f:
            cols=["booking_id","first_name","last_name","check_in","check_out","nights","price_usd","notes","source_sheet","source_row"]
            writer=csv.DictWriter(f,fieldnames=cols);writer.writeheader()
            writer.writerows({k:csv_literal(v) for k,v in row.items()} for row in model["accepted"])
        # 所有非数值自由文本封装在JSON内，避免直接成为CSV公式。
        with (temp/"2989_来源记录.csv").open("w",encoding="utf-8-sig",newline="") as f:
            writer=csv.writer(f);writer.writerow(["source_record_json"])
            writer.writerows([json.dumps(r,ensure_ascii=False,separators=(",",":"))] for r in model["ledger"])
        with (temp/"2989_期间汇总.csv").open("w",encoding="utf-8-sig",newline="") as f:
            writer=csv.DictWriter(f,fieldnames=["period","bookings","revenue_usd"]);writer.writeheader();writer.writerows(model["summary"])
        (temp/"2989_处理报告.json").write_text(json.dumps(model["report"],ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        if export_xlsx:
            subprocess.run([find_node(),str(Path(__file__).with_name("2989_生成工作簿.mjs")),"result","--model",str(model_path),"--output",str(temp/"2989_合并结果.xlsx"),"--previews",str(temp/"2989_工作簿预览")],check=True)
        if hashlib.sha256(source.read_bytes()).hexdigest()!=before:
            raise ValueError("处理期间源文件改变，取消发布结果")
        if output_dir.exists():
            raise ValueError("目标目录在处理期间出现，拒绝覆盖")
        temp.rename(output_dir)
    finally:
        if temp.exists():
            shutil.rmtree(temp)
    return model["report"]


def main():
    parser=argparse.ArgumentParser(description="2989 酒店月表保守合并；默认输出XLSX，--no-xlsx仅生成CSV/JSON")
    parser.add_argument("--source",required=True,type=Path)
    parser.add_argument("--output",required=True,type=Path)
    parser.add_argument("--bare-currency",choices=["USD"],default=None,help="仅在已确认裸金额均为USD时指定")
    parser.add_argument("--no-xlsx",action="store_true")
    args=parser.parse_args()
    try:
        print(json.dumps(process_workbook(args.source,args.output,args.bare_currency,not args.no_xlsx),ensure_ascii=False,indent=2))
    except Exception as exc:
        parser.exit(1,f"处理失败：{exc}\n")


if __name__=="__main__":
    main()
